"""Fase 3: coleta de histórico nos outros sistemas da Pro-Metal (somente leitura).

Cada fonte devolve {"fonte", "ok", "erro", "registros"}. Se uma tela mudar de layout ou o
sistema estiver fora do ar, a fonte volta com ok=False e o assistente avisa - nunca inventa dado.
Endereços ficam em instance/fontes.json (criado com os valores padrão na primeira execução).
"""
import json
import os
import re
import urllib.parse
import urllib.request
from datetime import date, timedelta
from html.parser import HTMLParser

PADRAO = {
    "cartao_amarelo": "http://192.168.1.150:5001",
    "refugo": "http://192.168.1.150:5003",
    "falhas": "http://192.168.1.150:5000",
    "materia_prima": "http://192.168.1.150:8765",
    "apqp": r"R:\APQP",
    "planilha_nc": r"P:\QUALIDADE\03. NÃO CONFORMIDADES\01. EXTERNAS"
                   r"\00. Gerenciamento de Reclamações 2026 - Externas - Rev 1.xlsx",
    "meses_historico": 12,
    # IA usada pelo assistente: "gemini" (chave em instance/gemini_key.txt)
    # ou "claude" (chave em instance/anthropic_key.txt)
    "provedor": "gemini",
    "modelo": "gemini-flash-latest",
    "modelos_reserva": ["gemini-flash-lite-latest"],
}
TIMEOUT = 15
MAX_REGISTROS = 40   # por fonte, para não estourar o pacote enviado ao Claude


def carregar_config(instance_path):
    caminho = os.path.join(instance_path, "fontes.json")
    if not os.path.exists(caminho):
        with open(caminho, "w", encoding="utf-8") as f:
            json.dump(PADRAO, f, ensure_ascii=False, indent=2)
    with open(caminho, encoding="utf-8") as f:
        return {**PADRAO, **json.load(f)}


# ---------------- leitura de tabelas HTML ----------------
class _Tabelas(HTMLParser):
    """Extrai todas as <table> como listas de linhas de texto (th e td)."""

    def __init__(self):
        super().__init__()
        self.tabelas, self._linha, self._celula = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.tabelas.append([])
        elif tag == "tr" and self.tabelas:
            self._linha = []
        elif tag in ("td", "th") and self._linha is not None:
            self._celula = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._celula is not None:
            self._linha.append(" ".join(" ".join(self._celula).split()))
            self._celula = None
        elif tag == "tr" and self._linha is not None:
            if self._linha:
                self.tabelas[-1].append(self._linha)
            self._linha = None

    def handle_data(self, data):
        if self._celula is not None:
            self._celula.append(data)


def _sem_tags(m):
    return " ".join(re.sub(r"<[^>]+>", " ", m.group(1)).split()) if m else ""


def _get(url, params=None):
    if params:
        url += "?" + urllib.parse.urlencode(params)
    with urllib.request.urlopen(url, timeout=TIMEOUT) as r:
        return r.read().decode("utf-8", errors="replace")


def _tabela_com(html, coluna):
    """Primeira tabela cujo cabeçalho contém `coluna`, como lista de dicts."""
    p = _Tabelas()
    p.feed(html)
    for t in p.tabelas:
        if t and coluna in t[0]:
            cab = t[0]
            return [dict(zip(cab, linha)) for linha in t[1:] if len(linha) == len(cab)]
    return None


def _resultado(fonte, registros=None, erro=None):
    registros = registros or []
    return {"fonte": fonte, "ok": erro is None, "erro": erro,
            "total": len(registros), "registros": registros[:MAX_REGISTROS]}


def _br(iso):
    """2026-03-09 -> 09/03/2026"""
    return "/".join(reversed(iso.split("-"))) if iso and iso[4:5] == "-" else iso


def _inicio(cfg):
    return date.today() - timedelta(days=30 * int(cfg["meses_historico"]))


# ---------------- fontes ----------------
def cartao_amarelo(pn, cfg):
    nome = "Cartão Amarelo/Vermelho"
    try:
        vistos, regs = set(), []
        for pagina in ("lotes", "concluidos"):
            linhas = _tabela_com(_get(f"{cfg['cartao_amarelo']}/{pagina}", {"codigo_produto": pn}), "Lote")
            if linhas is None:
                if "Nenhum lote" in _get(f"{cfg['cartao_amarelo']}/{pagina}", {"codigo_produto": pn}):
                    continue
                return _resultado(nome, erro=f"tabela não encontrada em /{pagina} (a tela mudou?)")
            for l in linhas:
                if l.get("Lote") in vistos:
                    continue
                vistos.add(l.get("Lote"))
                regs.append({**{k: v for k, v in l.items() if k not in ("Foto", "Ações") and v not in ("", "-")},
                             "link": f"{cfg['cartao_amarelo']}/{pagina}?" + urllib.parse.urlencode({"codigo_produto": pn})})
        return _resultado(nome, regs)
    except Exception as e:  # rede, timeout, HTML inesperado
        return _resultado(nome, erro=str(e))


def refugo(pn, cfg):
    nome = "Refugo"
    params = {"produto": pn, "inicio": _inicio(cfg).isoformat(), "fim": date.today().isoformat()}
    try:
        modos = json.loads(_get(f"{cfg['refugo']}/modo_falha", params)).get("modos", [])
        turnos = json.loads(_get(f"{cfg['refugo']}/produto_turno", params)).get("turnos", [])
        regs = [{"tipo": "modo de falha", **m} for m in modos] + [{"tipo": "turno", **t} for t in turnos]
        # datas, prensa e turno de cada um dos principais modos de falha (o Refugo não tem link
        # direto por peça: o link abre o dashboard, onde se filtra pela peça e pelo modo)
        for m in modos[:6]:
            d = json.loads(_get(f"{cfg['refugo']}/modo_falha", {**params, "modo": m["modo"]})).get("detalhe") or {}
            dia = d.get("diario") or {}
            datas = [f"{_br(x)} ({q} pç)" for x, q in zip(dia.get("dias", []), dia.get("qtde", []))]
            regs.append({"tipo": "datas do modo de falha", "modo": m["modo"], "qtde total": d.get("qtde"),
                         "lançamentos": d.get("lancamentos"),
                         "máquinas": ", ".join(f"{x['maquina']} ({x['qtde']} pç)" for x in d.get("maquinas", [])),
                         "turnos": ", ".join(f"{x['turno']} ({x['qtde']} pç)" for x in d.get("turnos", []) if x["qtde"]),
                         "primeira data": datas[0] if datas else "", "última data": datas[-1] if datas else "",
                         "datas": "; ".join(datas[-15:]), "link": f"{cfg['refugo']}/"})
        return _resultado(nome, regs)
    except Exception as e:
        return _resultado(nome, erro=str(e))


def falhas(pn, cfg):
    nome = "Dashboard de Falhas (alertas)"
    try:
        html = _get(f"{cfg['falhas']}/alertas", {"busca": pn})
        cartoes = html.split('class="alerta-card')[1:]
        if not cartoes:
            if "Nenhum alerta" in html:
                return _resultado(nome)
            return _resultado(nome, erro="cartões de alerta não encontrados (a tela mudou?)")
        regs = []
        for c in cartoes:
            campo = lambda cls: _sem_tags(re.search(rf'class="{cls}"[^>]*>(.*?)</(div|span)>', c, re.S))
            pdf = re.search(r"abrirAlerta\('([^']+)'", c)
            regs.append({"alerta": campo("titulo"), "peça": campo("peca-info"), "modo": campo("modo"),
                         "tipo": "Externo" if "tipo-Externos" in c[:60] else "Interno",
                         "data": _sem_tags(re.search(r'<div class="meta">.*?</span>\s*<span>(.*?)</span>', c, re.S)),
                         "link": cfg["falhas"] + urllib.parse.quote(pdf.group(1)) if pdf else
                                 f"{cfg['falhas']}/alertas?" + urllib.parse.urlencode({"busca": pn})})
        return _resultado(nome, regs)
    except Exception as e:
        return _resultado(nome, erro=str(e))


CAMPOS_BOBINA = ("id_bobina", "corrida", "fornecedor", "usina_origem", "codigo_mp", "data_entrada",
                 "data_certificado", "grau_aco", "espessura_mm", "largura_mm", "c", "mn", "si", "p", "s", "al",
                 "cr", "ni", "mo", "n", "ti", "cu", "nb", "le_mpa", "lr_mpa", "alongamento", "r90", "n90",
                 "dureza", "clc0", "norma", "status", "resultado")


def materia_prima(pn, cfg):
    """Certificados e CEP das bobinas (ensaios mecânicos e química) da MP usada na peça."""
    nome = "Matéria-prima (certificados e CEP)"
    try:
        # Código Pro-Metal: final "A" = peça final; "B", "B1", "C1"... = componentes estampados.
        # A MP entra nos componentes, então para a peça final buscamos todos os componentes dela
        # (mesmo código base). Ex.: FV10196A -> FV10196B1 e FV10196C1.
        pecas = json.loads(_get(f"{cfg['materia_prima']}/api/info")).get("pecas", [])
        base = re.sub(r"[A-Z]\d*$", "", pn)
        alvos = sorted({pn} & set(pecas)) + sorted(p for p in pecas if p != pn and re.sub(r"[A-Z]\d*$", "", p) == base)
        if not alvos:
            return _resultado(nome)
        if alvos != [pn]:
            nome += f" - peças consultadas: {', '.join(alvos)}"
        regs = []
        for peca in alvos[:4]:
            cep = json.loads(_get(f"{cfg['materia_prima']}/api/cep", {"peca": peca, "periodo": "tudo"}))
            regs += [{"tipo": "estatística CEP", "componente": peca, **e} for e in cep.get("estatisticas", [])]
            if cep.get("alerta_clc"):
                regs.append({"tipo": "alerta de estampagem (CLC0)", "componente": peca, "detalhe": cep["alerta_clc"]})
            bobinas = sorted(cep.get("bobinas", []), key=lambda b: b.get("id", 0), reverse=True)
            regs += [{"tipo": "bobina", "componente": peca,
                      **{k: b[k] for k in CAMPOS_BOBINA if b.get(k) not in (None, "")},
                      "link": f"{cfg['materia_prima']}/cep?" + urllib.parse.urlencode(
                          {"codigo_mp": b.get("codigo_mp") or "", "destaque": b.get("id")})}
                     for b in bobinas[:20 // len(alvos[:4])]]
            certs = json.loads(_get(f"{cfg['materia_prima']}/api/busca", {"peca": peca}))
            regs += [{"tipo": "certificado", "componente": peca,
                      **{k: c[k] for k in ("codigo_mp", "fornecedor", "material", "lote",
                                          "num_certificado", "nf", "data") if c.get(k)},
                      "link": f"{cfg['materia_prima']}/pdf?id={c['id']}"}
                     for c in certs[:5]]
        return _resultado(nome, regs)
    except Exception as e:
        return _resultado(nome, erro=str(e))


def _txt(v):
    if v is None:
        return ""
    return v.strftime("%d/%m/%Y") if hasattr(v, "strftime") else " ".join(str(v).split())


SEMAFORO = {"+G": "verde (+G)", "-G": "verde (-G)", "+Y": "amarelo (+Y)", "-Y": "amarelo (-Y)",
            "R": "VERMELHO - não realizado/atrasado (R)", "N/A": "não se aplica"}


def _cabecalho_dados(cab):
    """A aba Dados repete 'PRAZO' e 'Realizado' depois de cada etapa: renomeia para
    '<etapa> - prazo' / '<etapa> - realizado' e ignora as colunas de contagem após 'Observações'."""
    novo, etapa = [], ""
    for c in cab:
        if novo and novo[-1] == "Observações":
            novo += [""] * (len(cab) - len(novo))
            break
        if c in ("PRAZO", "Realizado"):
            novo.append(f"{etapa} - {c.lower()}")
        else:
            etapa = c if c not in novo else f"{c} (2ª vez)"   # 'CLIENTE INFORMADO' aparece 2x
            novo.append(etapa)
    return novo


def planilha_nc(pn, cfg):
    nome = "Planilha de NC externas"
    try:
        import openpyxl
        wb = openpyxl.load_workbook(cfg["planilha_nc"], read_only=True, data_only=True)
        alvo, regs = pn.upper(), []
        # aba Dados: cabeçalho na linha 5
        ws = wb["Dados"]
        linhas = ws.iter_rows(min_row=5, values_only=True)
        cab = _cabecalho_dados([_txt(c) for c in next(linhas)])
        for l in linhas:
            d = {cab[i]: SEMAFORO.get(_txt(v), _txt(v)) for i, v in enumerate(l)
                 if i < len(cab) and cab[i] and _txt(v)}
            codigos = " ".join(v for k, v in d.items() if "digo" in k).upper()
            if alvo in codigos:
                regs.append({"aba": "Dados", **d})
        # Planilha1: código -> modo de falha -> ação tomada
        if "Planilha1" in wb.sheetnames:
            for l in wb["Planilha1"].iter_rows(min_row=3, values_only=True):
                if l and l[0] and alvo in str(l[0]).upper():
                    regs.append({"aba": "Ações anteriores", "código": _txt(l[0]), "peça": _txt(l[1]),
                                 "data": _txt(l[2]), "modo de falha": _txt(l[3]), "ação": _txt(l[4])})
        wb.close()
        return _resultado(nome, regs)
    except Exception as e:
        return _resultado(nome, erro=str(e))


MAX_PDF_APQP = 6 * 1024 * 1024   # soma dos PDFs do APQP anexados (limite de envio à IA)


def _projetos_apqp(pn, cfg):
    """Linhas do índice geral APQP para a peça (coluna Totvs), da mais recente para a mais antiga."""
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(cfg["apqp"], "ORIGINAL - INDICE GERAL APQP.xlsx"),
                                read_only=True, data_only=True)
    base, achados = re.sub(r"[A-Z]\d*$", "", pn), []
    for l in wb["Índice"].iter_rows(min_row=8, values_only=True):
        totvs = str(l[7] or "").strip().upper()
        if totvs and re.sub(r"[A-Z]\d*$", "", totvs) == base and l[1] and l[3]:
            achados.append({"nº": l[1], "cliente": _txt(l[2]), "ano": _txt(l[3]), "responsável": _txt(l[4]),
                            "descrição": _txt(l[5]), "código debx": _txt(l[6]), "código totvs": totvs,
                            "projeto": _txt(l[8]), "aprovação ppap": _txt(l[9])})
    wb.close()
    return sorted(achados, key=lambda a: (a["código totvs"] != pn.upper(), -int(float(a["ano"] or 0))))


def _docs_apqp(pasta, debx):
    """PFMEA e plano de controle de produção mais recentes da peça dentro da pasta do projeto."""
    docs = []
    for raiz, _, arquivos in os.walk(pasta):
        if "obsolet" in raiz.lower():
            continue
        for a in arquivos:
            caminho = os.path.join(raiz, a)
            if not a.lower().endswith(".pdf") or debx.upper() not in caminho.upper():
                continue
            nome = a.lower()
            if "pfmea" in nome and "declara" not in nome:
                tipo = "PFMEA"
            elif "plano de controle" in caminho.lower() and "contenção" not in nome and "pré" not in raiz.lower():
                tipo = "Plano de controle"
            else:
                continue
            docs.append({"tipo": tipo, "caminho": caminho, "mtime": os.path.getmtime(caminho),
                         "tamanho": os.path.getsize(caminho), "pasta": raiz})
    escolhidos = []
    for tipo in ("PFMEA", "Plano de controle"):
        do_tipo = [d for d in docs if d["tipo"] == tipo]
        if not do_tipo:
            continue
        # a revisão mais recente é a pasta onde está o arquivo mais novo daquele tipo
        pasta_nova = max(do_tipo, key=lambda d: d["mtime"])["pasta"]
        escolhidos += sorted((d for d in do_tipo if d["pasta"] == pasta_nova), key=lambda d: d["caminho"])
    return escolhidos


def apqp(pn, cfg):
    """Projeto APQP da peça (R:\\APQP): dados do índice + PFMEA e plano de controle em PDF."""
    nome = "APQP (PFMEA e plano de controle)"
    try:
        projetos = _projetos_apqp(pn, cfg)
        if not projetos:
            return _resultado(nome)
        regs, anexos, total = [], [], 0
        for p in projetos:
            pasta_ano = os.path.join(cfg["apqp"], p["ano"])
            pastas = [d for d in (os.listdir(pasta_ano) if os.path.isdir(pasta_ano) else [])
                      if re.match(rf"0*{int(float(p['nº']))}(\s*-|\s|$)", d)]
            regs.append({"tipo": "projeto APQP", **p, "pasta": pastas[0] if pastas else "não encontrada"})
            if not pastas or anexos:          # documentos só do projeto mais recente que tiver
                continue
            for d in _docs_apqp(os.path.join(pasta_ano, pastas[0]), p["código debx"]):
                if total + d["tamanho"] > MAX_PDF_APQP:
                    regs.append({"tipo": "documento não anexado (limite de tamanho)", "arquivo": os.path.basename(d["caminho"])})
                    continue
                total += d["tamanho"]
                anexos.append({"tipo": d["tipo"], "caminho": d["caminho"]})
                regs.append({"tipo": f"{d['tipo']} anexado", "arquivo": os.path.basename(d["caminho"]),
                             "revisão (pasta)": os.path.relpath(d["pasta"], os.path.join(pasta_ano, pastas[0]))})
        r = _resultado(nome, regs)
        r["anexar"] = anexos
        return r
    except Exception as e:
        return _resultado(nome, erro=str(e))


FONTES = [cartao_amarelo, refugo, falhas, materia_prima, planilha_nc, apqp]


def coletar(pn, instance_path):
    cfg = carregar_config(instance_path)
    return [f(pn, cfg) for f in FONTES]
